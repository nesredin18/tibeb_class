import argparse
import html
import re
from difflib import SequenceMatcher
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook


KEHULUM_QUESTION_CANDIDATES = ["question", "question_text", "question_html"]
APP_QUESTION_CANDIDATES = ["question", "question_text", "question_html", "content"]
SUBJECT_CANDIDATES = ["subject"]
YEAR_CANDIDATES = ["year"]

ANSWER_COLUMNS = [
    "correct_letter",
    "correct_answer",
    "correct_answer_html",
    "correct_answer_image_urls",
    "explanation",
    "explanation_html",
    "explanation_image_urls",
]

META_COLUMNS = [
    "app_match_score",
    "app_matched_question",
    "app_matched_question_number",
    "app_matched_subject",
    "app_matched_year",
    "answer_option_match_score",
]


def read_excel_rows(path: Path) -> Tuple[List[str], List[Dict[str, str]]]:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    rows: List[Dict[str, str]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        item: Dict[str, str] = {}
        for index, header in enumerate(headers):
            value = row[index] if index < len(row) else ""
            item[header] = "" if value is None else str(value)
        rows.append(item)
    return headers, rows


def write_excel_rows(path: Path, headers: List[str], rows: List[Dict[str, str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Questions"
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    workbook.save(path)


def pick_first(row: Dict[str, str], candidates: List[str]) -> str:
    for name in candidates:
        value = row.get(name, "")
        if value:
            return value
    return ""


def html_to_text(value: str) -> str:
    if not value:
        return ""
    soup = BeautifulSoup(value, "html.parser")
    return soup.get_text(" ", strip=True)


def normalize_text(value: str) -> str:
    if not value:
        return ""
    value = html.unescape(value)
    value = html_to_text(value)
    value = value.lower()
    value = value.replace("\\text", " ")
    value = re.sub(r"\\[a-zA-Z]+", " ", value)
    value = re.sub(r"[_^{}[\]();,:.!?\"'`~|]", " ", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def normalize_year(value: str) -> str:
    match = re.search(r"(\d{4})", value or "")
    return match.group(1) if match else ""


def normalize_subject(value: str) -> str:
    value = (value or "").strip().lower()
    value = value.replace("_", " ")
    value = re.sub(r"\bnatural\b|\bsocial\b|\bscience\b", " ", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def similarity(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    return SequenceMatcher(None, a, b).ratio()


def kehulum_option_columns(row: Dict[str, str]) -> List[Tuple[str, str]]:
    options: List[Tuple[str, str]] = []
    for letter in ["A", "B", "C", "D", "E", "F"]:
        value = row.get(f"option_{letter.lower()}", "")
        html_value = row.get(f"option_{letter.lower()}_html", "")
        chosen = value or html_value
        if chosen:
            options.append((letter, chosen))
    return options


def find_best_kehulum_answer_letter(
    kehulum_row: Dict[str, str],
    app_match: Dict[str, str],
) -> Tuple[str, float]:
    app_answer_raw = app_match.get("correct_answer") or app_match.get("correct_answer_html") or ""
    app_answer_norm = normalize_text(app_answer_raw)
    if not app_answer_norm:
        return "", 0.0

    best_letter = ""
    best_score = 0.0
    for letter, option_value in kehulum_option_columns(kehulum_row):
        option_norm = normalize_text(option_value)
        score = similarity(app_answer_norm, option_norm)
        if score > best_score:
            best_score = score
            best_letter = letter
    return best_letter, best_score


def build_app_index(app_rows: List[Dict[str, str]]) -> List[Dict[str, str]]:
    indexed: List[Dict[str, str]] = []
    for row in app_rows:
        question = pick_first(row, APP_QUESTION_CANDIDATES)
        indexed.append(
            {
                **row,
                "__question_raw": question,
                "__question_norm": normalize_text(question),
                "__subject_norm": normalize_subject(pick_first(row, SUBJECT_CANDIDATES)),
                "__year_norm": normalize_year(pick_first(row, YEAR_CANDIDATES)),
            }
        )
    return indexed


def find_best_match(
    kehulum_row: Dict[str, str],
    app_index: List[Dict[str, str]],
    threshold: float,
) -> Tuple[Optional[Dict[str, str]], float]:
    kehulum_question = pick_first(kehulum_row, KEHULUM_QUESTION_CANDIDATES)
    kehulum_norm = normalize_text(kehulum_question)
    kehulum_subject = normalize_subject(pick_first(kehulum_row, SUBJECT_CANDIDATES))
    kehulum_year = normalize_year(pick_first(kehulum_row, YEAR_CANDIDATES))

    best_row: Optional[Dict[str, str]] = None
    best_score = 0.0

    candidates = app_index
    filtered = [
        row
        for row in app_index
        if (not kehulum_year or row["__year_norm"] == kehulum_year)
        and (not kehulum_subject or not row["__subject_norm"] or row["__subject_norm"] == kehulum_subject)
    ]
    if filtered:
        candidates = filtered

    for app_row in candidates:
        score = similarity(kehulum_norm, app_row["__question_norm"])
        if score > best_score:
            best_score = score
            best_row = app_row

    if best_score < threshold:
        return None, best_score
    return best_row, best_score


def merge_rows(
    kehulum_rows: List[Dict[str, str]],
    app_rows: List[Dict[str, str]],
    threshold: float,
) -> List[Dict[str, str]]:
    app_index = build_app_index(app_rows)
    merged: List[Dict[str, str]] = []

    for kehulum_row in kehulum_rows:
        output_row = dict(kehulum_row)
        match, score = find_best_match(kehulum_row, app_index, threshold)

        for column in ANSWER_COLUMNS + META_COLUMNS:
            output_row.setdefault(column, "")

        if match:
            for column in ANSWER_COLUMNS:
                output_row[column] = match.get(column, "")

            resolved_letter, option_match_score = find_best_kehulum_answer_letter(output_row, match)
            if resolved_letter:
                output_row["correct_letter"] = resolved_letter
            output_row["answer_option_match_score"] = f"{option_match_score:.4f}"

            output_row["app_match_score"] = f"{score:.4f}"
            output_row["app_matched_question"] = match.get("__question_raw", "")
            output_row["app_matched_question_number"] = match.get("question_number", "")
            output_row["app_matched_subject"] = match.get("subject", "")
            output_row["app_matched_year"] = match.get("year", "")
        else:
            output_row["app_match_score"] = f"{score:.4f}"
            output_row["answer_option_match_score"] = "0.0000"

        merged.append(output_row)

    return merged


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Match Kehulum questions to app-exported questions and copy answers into a new Excel file."
    )
    parser.add_argument("--kehulum", required=True, help="Path to the Kehulum Excel file.")
    parser.add_argument("--app", required=True, help="Path to the app export Excel file.")
    parser.add_argument("--output", required=True, help="Path to the merged Excel output file.")
    parser.add_argument(
        "--threshold",
        type=float,
        default=0.72,
        help="Minimum text similarity required to accept a match. Default: 0.72",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_arguments()
    kehulum_path = Path(args.kehulum).resolve()
    app_path = Path(args.app).resolve()
    output_path = Path(args.output).resolve()

    kehulum_headers, kehulum_rows = read_excel_rows(kehulum_path)
    _, app_rows = read_excel_rows(app_path)

    merged_rows = merge_rows(kehulum_rows, app_rows, args.threshold)

    output_headers = list(kehulum_headers)
    for column in ANSWER_COLUMNS + META_COLUMNS:
        if column not in output_headers:
            output_headers.append(column)

    write_excel_rows(output_path, output_headers, merged_rows)
    matched_count = sum(1 for row in merged_rows if row.get("correct_letter"))
    print(f"Done. Matched {matched_count} of {len(merged_rows)} Kehulum rows.")
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    main()
